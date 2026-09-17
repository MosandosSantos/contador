import csv
import io
import re
from collections import defaultdict
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from psycopg2.extras import RealDictCursor, Json
from core.db import connection

CAMPOS=['linha_id','documento_id','empresa_id','data','conta_id','debito_centavos','credito_centavos','tipo','historico','origem_id','centro_id','atividade_caixa','rubrica_caixa']
TIPOS={'abertura','receita','recebimento','pagamento','competencia','depreciacao','aporte','distribuicao','encerramento','reserva','investimento'}


def celula(valor):
    """Célula de CSV/XLSX como texto limpo; datas viram AAAA-MM-DD."""
    if valor is None:
        return ''
    if isinstance(valor, datetime):
        return valor.date().isoformat()
    if isinstance(valor, date):
        return valor.isoformat()
    if isinstance(valor, bool):
        return '1' if valor else '0'
    if isinstance(valor, float):
        return str(int(valor)) if valor.is_integer() else repr(valor)
    return str(valor).strip()


def decimal_celula(valor):
    """Célula como Decimal ('1.000,50', '1000.50', número do Excel). None se vazia."""
    if valor is None:
        return None
    if isinstance(valor, (int, float)) and not isinstance(valor, bool):
        return Decimal(str(valor))
    texto = str(valor).strip().replace(' ', '')
    if not texto:
        return None
    if '.' in texto and ',' in texto:
        texto = texto.replace('.', '').replace(',', '.')
    elif ',' in texto:
        texto = texto.replace(',', '.')
    try:
        return Decimal(texto)
    except InvalidOperation:
        raise ValueError('Valor numérico inválido: ' + str(valor)[:40])


def _cabecalho(valor):
    """Cabeçalho de coluna em minúsculas, sem espaços (vira _) nem acentos."""
    texto = celula(valor).lower()
    from unicodedata import normalize
    return re.sub(r'\s+', '_', normalize('NFKD', texto).encode('ascii', 'ignore').decode())


def ler_tabela(bruto: bytes, nome: str):
    """Lê CSV (UTF-8/Latin-1, vírgula ou ponto e vírgula) ou a primeira aba de XLSX.

    Devolve dicionários por linha com cabeçalho em minúsculas, sem espaços/acentos.
    """
    if nome.lower().endswith('.csv'):
        for encoding in ('utf-8-sig', 'latin-1'):
            try:
                texto = bruto.decode(encoding)
                break
            except UnicodeDecodeError:
                continue
        else:
            raise ValueError('Não consegui decodificar o CSV.')
        amostra = texto[:2048]
        separador = ';' if amostra.count(';') > amostra.count(',') else ','
        leitor = csv.DictReader(io.StringIO(texto), delimiter=separador)
        if not leitor.fieldnames or len(set(leitor.fieldnames)) != len(leitor.fieldnames):
            raise ValueError('Cabeçalho inválido ou repetido.')
        linhas = []
        for linha in leitor:
            if any((v or '').strip() for v in linha.values() if isinstance(v, str)):
                linhas.append(dict((_cabecalho(k), v) for k, v in linha.items() if k is not None))
        return linhas
    if nome.lower().endswith(('.xlsx', '.xlsm')):
        from openpyxl import load_workbook
        try:
            livro = load_workbook(io.BytesIO(bruto), read_only=True, data_only=True)
        except Exception as erro:
            raise ValueError('Não consegui ler a planilha: ' + str(erro)) from erro
        try:
            aba = livro[livro.sheetnames[0]]
            brutas = list(aba.iter_rows(values_only=True))
        finally:
            livro.close()
        if not brutas:
            return []
        cabecalho = [_cabecalho(c) for c in brutas[0]]
        if not any(cabecalho) or len(set(cabecalho)) != len(cabecalho):
            raise ValueError('Cabeçalho inválido ou repetido.')
        return [dict(zip(cabecalho, linha)) for linha in brutas[1:]
                if any(c is not None and str(c).strip() for c in linha)]
    raise ValueError('Formato não suportado: use CSV ou XLSX.')

def ler_csv(raw):
    try: text=raw.decode('utf-8-sig')
    except UnicodeDecodeError: raise ValueError('Use CSV UTF-8.')
    reader=csv.DictReader(io.StringIO(text))
    if not reader.fieldnames or len(set(reader.fieldnames))!=len(reader.fieldnames): raise ValueError('Cabeçalho inválido ou repetido.')
    rows=list(reader)
    if not rows or len(rows)>5000: raise ValueError('Envie de 1 a 5.000 linhas por arquivo.')
    if any(None in row or any(v is None for v in row.values()) for row in rows): raise ValueError('CSV com quantidade de colunas inconsistente.')
    return rows

def validar_lancamentos(rows, contas):
    ids=set();docs=defaultdict(list)
    for row in rows:
        if set(row)!=set(CAMPOS): raise ValueError('Colunas diferentes do modelo. Baixe o CSV de exemplo.')
        if row['linha_id'] in ids: raise ValueError('Identificador de linha duplicado no arquivo.')
        ids.add(row['linha_id'])
        if any(len(str(v))>300 for v in row.values()): raise ValueError('Campo maior que 300 caracteres.')
        if not row['linha_id'] or not row['documento_id']: raise ValueError('Informe identificadores de linha e documento.')
        if row['conta_id'] not in contas: raise ValueError('Conta não mapeada: '+row['conta_id'])
        try: data_lancamento = date.fromisoformat(row['data'])
        except (TypeError, ValueError): raise ValueError('Data inválida.')
        if not 1900 <= data_lancamento.year <= 2199: raise ValueError('Ano fora de 1900–2199.')
        if row['empresa_id']!='EMPRESA_WORKSHOP_01': raise ValueError('Empresa diferente do caso fictício.')
        if row['tipo'] not in TIPOS: raise ValueError('Tipo de documento inválido.')
        for k in ('debito_centavos','credito_centavos'):
            if not re.fullmatch(r'\d{1,12}',str(row[k])): raise ValueError('Valor monetário deve ser inteiro não negativo, em centavos.')
            row[k]=int(row[k])
        if row['debito_centavos'] and row['credito_centavos']: raise ValueError('Uma linha não pode ter débito e crédito juntos.')
        if not row['debito_centavos'] and not row['credito_centavos']: raise ValueError('Linha sem débito nem crédito não é um lançamento.')
        if row['conta_id']=='1.1.1' and row['tipo']!='abertura' and row['atividade_caixa'] not in ('operacional','investimento','financiamento'):
            raise ValueError('Movimento bancário sem atividade da DFC.')
        docs[row['documento_id']].append(row)
    for doc,lines in docs.items():
        if len({(x['data'],x['tipo']) for x in lines})!=1: raise ValueError('Documento com datas ou tipos diferentes: '+doc)
        if sum(x['debito_centavos']-x['credito_centavos'] for x in lines)!=0: raise ValueError('Documento desbalanceado: '+doc)
    return rows

def importar(rows):
    with connection() as conn, conn.cursor(cursor_factory=RealDictCursor) as cur:
        cur.execute('SELECT pg_advisory_xact_lock(20260911)')
        cur.execute('SELECT conta_id FROM contas');contas={x['conta_id'] for x in cur.fetchall()}
        validar_lancamentos(rows,contas)
        cur.execute('SELECT * FROM lancamentos');existing={x['linha_id']:dict(x) for x in cur.fetchall()}
        documents={x['documento_id'] for x in existing.values()}
        novas=[]
        for x in rows:
            antigo=existing.get(x['linha_id'])
            if antigo:
                antigo['data']=antigo['data'].isoformat()
                if any(antigo[k]!=x[k] for k in CAMPOS):raise ValueError('Identificador existente com conteúdo diferente: '+x['linha_id'])
            else:
                if x['documento_id'] in documents: raise ValueError('Documento existente: não é permitido acrescentar linhas por importação.')
                if x['tipo']=='abertura':raise ValueError('A abertura já existe. Uma segunda abertura não é permitida.')
                novas.append(x)
        for x in novas:
            cur.execute('INSERT INTO lancamentos VALUES ('+','.join(['%s']*len(CAMPOS))+')',tuple(x[k] for k in CAMPOS))
        cur.execute('INSERT INTO auditoria(acao,antes,depois) VALUES(%s,%s,%s)',('Importação',Json({'linhas':len(existing)}),Json({'linhas':len(existing)+len(novas),'novas':len(novas),'ignoradas':len(rows)-len(novas)})))
    return len(novas),len(rows)-len(novas)

def validar_extrato(rows):
    ids=set()
    for x in rows:
        if not {'registro_id','referencia','valor_centavos','data'}<=set(x):raise ValueError('Use o modelo de extrato.')
        if not x['registro_id'] or x['registro_id'] in ids:raise ValueError('Identificador de registro duplicado.')
        ids.add(x['registro_id'])
        if not x['referencia'] or len(x['referencia'])>100:raise ValueError('Referência inválida.')
        if not re.fullmatch(r'-?\d{1,12}',str(x['valor_centavos'])):raise ValueError('Valor deve estar em centavos.')
        x['valor_centavos']=int(x['valor_centavos'])
        try: date.fromisoformat(x['data'])
        except ValueError:raise ValueError('Data inválida no extrato.')
    return rows
